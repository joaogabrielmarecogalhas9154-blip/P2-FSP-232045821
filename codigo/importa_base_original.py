"""
Importa a base do TRABALHO ANTERIOR (P1, João Paulo Cortes da Cunha Dutra, 241021838) para formato
longo/tabular, sem alterar os arquivos originais, e valida o extrator de rubricas contra ela.

Fonte: https://github.com/Jhons777/P1-Finan-as-do-Setor-Publico---Jo-o-Paulo-Cortes-Da-Cunha-Dutra-241021838
Commit usado: 4ae3ff41d0895097b752359aad152a4cac551f33 (08/05/2026)

Uso:  python importa_base_original.py data/original/P1_anterior
Saída (em data/original/):
  contas_receita_original.csv   todas as contas de receita das abas brutas (mun, ano, código, rótulo, valor)
  painel_original.csv           painel nominal da P1 (receitas, despesas, funções, população, PIB)
  validacao_extrator.csv        extrator novo × colunas da P1 (diferença absoluta e relativa)
  achados_auditoria.csv         inconsistências detectadas automaticamente
"""
import os
import re
import sys
import numpy as np
import pandas as pd
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from rubricas_receita import extrai, separa_codigo  # noqa: E402

ORIG = sys.argv[1] if len(sys.argv) > 1 else os.path.join("data", "original", "P1_anterior")
OUT = os.path.join("data", "original")
os.makedirs(OUT, exist_ok=True)
COD = {"Barreiras": 2903201, "Camaçari": 2905701}
ABA = re.compile(r"^\s*(Barreiras|Camaçari)\s*-\s*BA\s*\((\d{4})\)\s*$")


def mun_nome(txt):
    return "Barreiras" if str(txt).startswith("Barreiras") else "Camaçari"


def contas_brutas(arquivo):
    wb = openpyxl.load_workbook(arquivo, read_only=True, data_only=True)
    linhas = []
    for nome in wb.sheetnames:
        m = ABA.match(nome)
        if not m:
            continue
        mun, ano = m.group(1), int(m.group(2))
        for r in wb[nome].iter_rows(values_only=True):
            if len(r) < 3 or r[1] is None:
                continue
            codigo, rotulo = separa_codigo(str(r[1]))
            if codigo is None and not str(r[1]).upper().startswith(("TOTAL", "RECEITAS (EXCETO")):
                continue
            valor = pd.to_numeric(r[2], errors="coerce")
            linhas.append({"municipio": mun, "cod_ibge": COD[mun], "ano": ano,
                           "codigo": codigo, "rotulo": rotulo if codigo else str(r[1]), "valor": valor})
    return pd.DataFrame(linhas)


def main():
    rd = os.path.join(ORIG, "Dados Brutos")
    rec = contas_brutas(os.path.join(rd, "Receitas orçamentarias", "Base_receitas (Principal).xlsx"))
    rec.to_csv(os.path.join(OUT, "contas_receita_original.csv"), index=False, encoding="utf-8-sig")

    # ---------------- painel da P1 (valores nominais, como publicados)
    br = pd.read_excel(os.path.join(rd, "Receitas orçamentarias", "Base_receitas (Principal).xlsx"), "Base_Receitas")
    bd = pd.read_excel(os.path.join(rd, "Despesas orçamentarias", "base_despesas (principal).xlsx"), "Base_Despesas")
    bf = pd.read_excel(os.path.join(rd, "Despesas por função", "base_despesas_funcao_2014_2023.xlsx"), "Base_Despesas_Funcao")
    bp = pd.read_excel(os.path.join(rd, "Base_Populacao_Municipal.xlsx"), "Base_Populacao")
    bpib = pd.read_excel(os.path.join(rd, "Base_PIB_corrigida.xlsx"), "Base_PIB")
    for d in (br, bd, bf, bp, bpib):
        d["municipio"] = d["Município"].map(mun_nome)
    keep = lambda d, cols: d[["Ano", "municipio"] + cols]  # noqa: E731
    painel = keep(br, ["Receita_Total", "Receita_Tributaria", "IPTU", "ISS", "ITBI", "Transferencias", "FPM",
                       "ICMS", "FUNDEB", "Receita_Capital"])
    painel = painel.merge(keep(bd, [c for c in bd.columns if c.startswith(("Despesa", "Pessoal", "Juros", "Outras",
                                                                         "Investimentos", "Inversoes", "Amortizacao"))]),
                          on=["Ano", "municipio"], how="outer")
    painel = painel.merge(keep(bf, ["Educacao", "Saude", "Saneamento", "Urbanismo", "Assistencia_Social",
                                    "Administracao"]).add_prefix("Funcao_").rename(
        columns={"Funcao_Ano": "Ano", "Funcao_municipio": "municipio"}), on=["Ano", "municipio"], how="outer")
    painel = painel.merge(keep(bp, ["População"]), on=["Ano", "municipio"], how="outer")
    painel = painel.merge(keep(bpib, ["PIB", "PIB_per_capita"]), on=["Ano", "municipio"], how="outer")
    painel = painel.rename(columns={"Ano": "ano"}).sort_values(["municipio", "ano"])
    painel["cod_ibge"] = painel["municipio"].map(COD)
    painel.to_csv(os.path.join(OUT, "painel_original.csv"), index=False, encoding="utf-8-sig")

    # ---------------- validação do extrator
    comp = {"IPTU": "iptu", "ISS": "iss", "ITBI": "itbi", "FPM": "fpm", "ICMS": "icms_cota", "FUNDEB": "fundeb"}
    val, usados_log = [], []
    for (mun, ano), g in rec[rec.codigo.notna()].groupby(["municipio", "ano"]):
        v, usados = extrai(g)
        p1 = br[(br.municipio == mun) & (br.Ano == ano)]
        for col, conc in comp.items():
            a = float(p1[col].iloc[0]) if len(p1) else np.nan
            val.append({"municipio": mun, "ano": ano, "variavel": conc, "valor_P1": a, "valor_extrator": v[conc],
                        "dif_abs": v[conc] - a, "dif_rel": (v[conc] - a) / a if a else np.nan,
                        "contas_usadas": ";".join(usados[conc])})
        for k in ["itr", "ipva_cota", "ipi_exp_cota", "iof_ouro", "outras_multigov", "sus", "transf_incondicionais"]:
            val.append({"municipio": mun, "ano": ano, "variavel": k, "valor_P1": np.nan, "valor_extrator": v[k],
                        "dif_abs": np.nan, "dif_rel": np.nan, "contas_usadas": ";".join(usados.get(k, []))})
    val = pd.DataFrame(val)
    val.to_csv(os.path.join(OUT, "validacao_extrator.csv"), index=False, encoding="utf-8-sig")

    # ---------------- achados automáticos
    ach = []
    for c in [c for c in painel.columns if c.startswith("Funcao_")]:
        for _, r in painel[painel[c] == 0].iterrows():
            ach.append({"tipo": "zero_no_lugar_de_ausente", "municipio": r.municipio, "ano": r.ano, "variavel": c,
                        "detalhe": ("P1 gravou 0, mas a função não consta na aba bruta do ano. Tratado como ausente (NaN). "
                                   "Saúde/Administração: certamente lacuna de coleta; Saneamento: pode ser gasto nulo "
                                   "(a DCA omite funções sem empenho) - conferir na nova coleta")})
    for _, r in painel[painel["População"].isna()].iterrows():
        ach.append({"tipo": "populacao_ausente", "municipio": r.municipio, "ano": r.ano, "variavel": "População",
                    "detalhe": "P1 não preencheu; per capita do ano ficou vazio"})
    fp = val[(val.variavel == "fpm") & (val.dif_rel.abs() >= 1e-3)]
    for _, r in fp.iterrows():
        ach.append({"tipo": "definicao_inconsistente", "municipio": r.municipio, "ano": r.ano, "variavel": "FPM",
                    "detalhe": (f"P1 = R$ {r.valor_P1:,.2f} (só cota mensal); com as parcelas de 1% = "
                                f"R$ {r.valor_extrator:,.2f}. A partir de 2018 a P1 soma as parcelas: quebra de definição")})
    om = val[(val.variavel == "outras_multigov") & (val.valor_extrator > 0)]
    for _, r in om.iterrows():
        ach.append({"tipo": "transferencia_extraordinaria", "municipio": r.municipio, "ano": r.ano,
                    "variavel": "Outras Transferências Multigovernamentais",
                    "detalhe": f"R$ {r.valor_extrator:,.2f} incluídos em Receita_Total/Transferencias da P1"})
    for mun, g in painel.groupby("municipio"):
        g = g.sort_values("ano")
        var = g["Receita_Total"].pct_change()
        for ano, x in zip(g.ano, var):
            if pd.notna(x) and abs(x) > 0.25:
                ach.append({"tipo": "salto_receita_total", "municipio": mun, "ano": ano, "variavel": "Receita_Total",
                            "detalhe": f"variação nominal de {x:+.1%} em relação ao ano anterior"})
    pd.DataFrame(ach).to_csv(os.path.join(OUT, "achados_auditoria.csv"), index=False, encoding="utf-8-sig")

    v6 = val[val.variavel.isin(comp.values())]
    print("Contas de receita lidas:", len(rec))
    print("Validação (|dif_rel| < 0,1%):", (v6.dif_rel.abs() < 1e-3).sum(), "de", v6.dif_rel.notna().sum())
    print(v6[v6.dif_rel.abs() >= 1e-3][["municipio", "ano", "variavel", "valor_P1", "valor_extrator", "dif_rel"]]
          .to_string(index=False))
    print("Achados de auditoria:", len(ach))


if __name__ == "__main__":
    main()
